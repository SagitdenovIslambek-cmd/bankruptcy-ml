"""Генерация отчёта (docx, 5–7 стр.) из artifacts/results.json и графиков. Все числа берутся из результатов запуска."""
from __future__ import annotations

import json

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .data import ROOT
from .features import NEW_FEATURES, feature_rationales

ART, FIG = ROOT / "artifacts", ROOT / "reports" / "figures"
FONT = "Times New Roman"


def f(x, d=3):
    return f"{x:.{d}f}".replace(".", ",")


def pc(x, d=1):
    return f"{x * 100:.{d}f}".replace(".", ",") + "%"


def _shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def _cell_margins(table, top=20, bottom=20, left=60, right=60):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tblPr.append(m)


class Doc:
    def __init__(self):
        self.d = Document()
        s = self.d.sections[0]
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin, s.bottom_margin = Cm(1.8), Cm(1.8)
        st = self.d.styles["Normal"]
        st.font.name, st.font.size = FONT, Pt(10)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.line_spacing = 1.05
        for name, size in (("Heading 1", 12.5), ("Heading 2", 11)):
            h = self.d.styles[name]
            h.font.name, h.font.size, h.font.bold = FONT, Pt(size), True
            h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(8), Pt(3)

    def h(self, t, lvl=1):
        self.d.add_heading(t, lvl)

    def p(self, text, bold_prefix=None, size=None, italic=False, align=None):
        para = self.d.add_paragraph()
        if bold_prefix:
            r = para.add_run(bold_prefix); r.bold = True
            if size: r.font.size = Pt(size)
        r = para.add_run(text); r.italic = italic
        if size: r.font.size = Pt(size)
        if align == "j":
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return para

    def bullets(self, items, size=None):
        for it in items:
            para = self.d.add_paragraph(style="List Bullet")
            para.paragraph_format.space_after = Pt(1)
            if isinstance(it, tuple):
                r = para.add_run(it[0]); r.bold = True
                r2 = para.add_run(it[1])
                if size: r.font.size = r2.font.size = Pt(size)
            else:
                r = para.add_run(it)
                if size: r.font.size = Pt(size)

    def table(self, header, rows, widths=None, size=8.5, align_right_from=1, bold_last=False):
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _cell_margins(t)
        for i, h_ in enumerate(header):
            c = t.rows[0].cells[i]; c.text = ""
            r = c.paragraphs[0].add_run(str(h_)); r.bold = True; r.font.size = Pt(size)
            _shade(c, "DCE6F1")
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                para = cells[i].paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                r = para.add_run(str(v)); r.font.size = Pt(size)
                if bold_last and ri == len(rows) - 1: r.bold = True
                if i >= align_right_from and align_right_from >= 0:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if widths:
            t.autofit = False
            lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(lay)
            for i, w in enumerate(widths):
                t.columns[i].width = Cm(w)
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self.d.add_paragraph().paragraph_format.space_after = Pt(1)
        return t

    def image(self, path, width_cm, caption=None):
        para = self.d.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(str(path), width=Cm(width_cm))
        para.paragraph_format.space_after = Pt(0)
        if caption:
            c = self.d.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = c.add_run(caption); r.italic = True; r.font.size = Pt(8.5)

    def two_images(self, a, b, wa, wb, caption=None):
        t = self.d.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, img, w in zip(t.rows[0].cells, (a, b), (wa, wb)):
            para = cell.paragraphs[0]; para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.add_run().add_picture(str(img), width=Cm(w))
        if caption:
            c = self.d.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = c.add_run(caption); r.italic = True; r.font.size = Pt(8.5)


def _short(rat: str) -> str:
    return rat.split(":")[0].strip() if ":" in rat[:90] else rat.split(".")[0]


def build_report(out=None) -> str:
    R = json.load(open(ART / "results.json", encoding="utf-8"))
    s1, s2, s3, s4 = R["step1"], R["step2"], R["step3"], R["step4"]
    eng = s4["engine"]
    prelim = eng != "lightgbm" or s4["shap_method"] != "shap"
    eng_name = "LightGBM" if eng == "lightgbm" else "HistGradientBoosting (sklearn)"
    shap_name = "SHAP (TreeExplainer)" if s4["shap_method"] == "shap" else "вклады признаков методом замещения медианой"
    B, Bs = s4["boost"], s4["base"]
    fin = s4["finance"]
    sp = {r["sample"]: r for r in s1["split"]}
    tr, va, te = sp["train"], sp["valid"], sp["test (OOT)"]
    D = Doc()

    # ---------- титул ----------
    t = D.d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Прогноз банкротства компаний: скоринговая карта и градиентный бустинг"); r.bold = True; r.font.size = Pt(15)
    t = D.d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.add_run("Практическое задание «ML-моделирование в финансах», вариант 5. Данные: Kaggle «Company Bankruptcy Prediction» (Тайвань)").font.size = Pt(9.5)
    if prelim:
        w = D.d.add_paragraph(); w.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = w.add_run("ПРЕДВАРИТЕЛЬНАЯ ВЕРСИЯ: расчёт выполнен на запасном движке (" + eng_name + "; объяснения — " + shap_name +
                      "). Финальные числа получаются командой python run_all.py при установленных LightGBM и SHAP.")
        r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)

    # ---------- 1. постановка ----------
    D.h("1. Постановка задачи")
    D.p(f"Требуется по годовой финансовой отчётности компании оценить вероятность банкротства в течение года (PD) и "
        f"превратить её в решение «одобрить / отказать». Целевая переменная — Bankrupt? (0/1). Ключевые метрики варианта: "
        f"Gini и полнота (доля пойманных банкротов) при фиксированной доле проверяемых компаний; затем качество переводится "
        f"в прибыль портфеля при LGD = 45%. Сравниваются интерпретируемая базовая модель (WoE-скоринговая карта + логистическая "
        f"регрессия) и градиентный бустинг ({eng_name}); решения бустинга объясняются через {shap_name}.", align="j")

    # ---------- 2. данные ----------
    D.h("2. Данные, разбиение по времени, проверка утечки")
    D.p(f"Выборка: {s1['n_rows']} компаний, {s1['n_raw_features']} признаков (коэффициенты отчётности: рентабельность, ликвидность, "
        f"долговая нагрузка, оборачиваемость, денежные потоки, рост), пропусков {s1['missing_total']}; банкротств "
        f"{s1['bankrupt_total']} ({pc(s1['bankrupt_total'] / s1['n_rows'])}). Тип всех признаков — число; доля пропусков и экономический "
        f"смысл каждого признака приведены в artifacts/features_description.csv и ноутбуке 01.", align="j")
    D.p("В датасете нет колонки с датой отчёта (данные Taiwan Economic Journal, 1999–2009, записаны хронологически). "
        "Поэтому порядок строк использован как прокси времени: подтверждение — доля банкротств в целом резко снижается по порядку строк "
        f"({pc(min(s1['block_rates'][:4]), 0)}–{pc(max(s1['block_rates'][:4]), 0)} в первых четырёх десятых частях и "
        f"{pc(min(s1['block_rates'][4:]), 0)}–{pc(max(s1['block_rates'][4:]), 0)} в последних шести, рис. 1). "
        "Это главное ограничение работы: календарные даты недоступны, поэтому «out-of-time» понимается как «более поздние строки».",
        bold_prefix="Дата наблюдения. ", align="j")
    D.table(["Выборка", "Строки (порядок)", "Наблюдений", "Банкротств", "Доля банкротств"],
            [[n, f"{r['row_from']}–{r['row_to']}", r["n"], r["bankrupt"], pc(r["bankrupt_rate"], 2)]
             for n, r in (("train (70%)", tr), ("validation (15%)", va), ("test, out-of-time (15%)", te))],
            widths=[4.2, 3.5, 3, 3, 3.3])
    d = s1["split_demo"]
    D.p(f"Разбиение выполнено по времени: train — первые 70%, validation — следующие 15%, test — последние 15%; test не использовался "
        f"до этапа 4 и открывается один раз в коде. Для сравнения, одна и та же простая модель даёт Gini {f(d['gini_valid_time_split'])} "
        f"на validation при разбиении по времени и {f(d['gini_valid_random_split'])} при случайном: при дрейфе доли банкротов оценка на "
        f"перемешанных периодах некорректна, поэтому случайное разбиение не применялось. Все статистики (границы винзоризации, пороги "
        f"флага стресса, WoE-корзины, медианы) считаются только на train.", bold_prefix="Разбиение. ", align="j")
    rem = pd.read_csv(ART / "removed_features.csv")
    const = rem[rem.reason == "constant"].feature.tolist()
    dup = rem[rem.reason == "duplicate"]
    leak = rem[rem.reason == "leak_suspect"]
    D.p(f"Проверена по трём направлениям. (1) По смыслу: все 95 признаков — коэффициенты отчётности за период до банкротства; признаков, "
        f"становящихся известными только после события (суммы просрочки, «будущие» цены), нет. (2) Статистически: максимальный "
        f"одномерный Gini признака на train — {f(s1['max_univariate_gini'])}, подозрительно высоких (> 0,90) значений "
        f"{'нет' if leak.empty else 'есть: ' + ', '.join(leak.feature)}. (3) Избыточность: удалено {len(rem)} признаков — "
        f"константа {', '.join(const) if const else '—'} и {len(dup)} дублей (|ρ Спирмена| ≥ 0,98; оставлен признак с большим Gini на train). "
        f"Каждый удалённый признак перечислен в таблице.", bold_prefix="Утечка данных. ", align="j")
    rows = [[r.feature, "константа" if r.reason == "constant" else ("дубль" if r.reason == "duplicate" else "подозрение на утечку"),
             (r.detail + (f"; оставлен: {r.partner}" if isinstance(r.partner, str) and r.partner else ""))] for r in rem.itertuples()]
    D.table(["Удалённый признак", "Причина", "Обоснование"], rows, widths=[6.2, 2.6, 8.2], size=7.5, align_right_from=-1)
    D.two_images(FIG / "fig1_drift.png", FIG / "fig_calibration_test.png", 8.6, 5.6,
                 "Рис. 1. Дрейф доли банкротств по порядку строк (слева); калибровка на test (справа)")

    # ---------- 3. признаки ----------
    D.h("3. Признаки")
    D.p(f"После аудита оставлено {s1['n_features_after_audit']} исходных признаков; добавлено {s1['n_new_features']} новых "
        f"(финальная матрица — {s3['n_features_model']} столбцов). Каждый признак имеет экономическое обоснование; значения винзоризованы "
        f"по 1%/99% квантилям train, деления защищены сдвигом знаменателя.", align="j")
    rat = feature_rationales()
    sh_new = s4.get("shap_top_new", {})
    D.table(["Признак", "Экономический смысл"],
            [[n, rat[n].split(". ")[0].rstrip(".") + "."] for n in NEW_FEATURES], widths=[4.2, 12.8], size=7.5, align_right_from=-1)
    g_raw, g_all = s3["gini_valid_raw_features_only"], s3["valid_raw"]["gini"]
    D.p(f"Вклад новых признаков: бустинг с теми же параметрами на одних исходных признаках даёт Gini {f(g_raw)} на validation, "
        f"со всеми — {f(g_all)}" + (" (разница в пределах шума: на validation всего " + str(va["bankrupt"]) + " банкротов). " if abs(g_all - g_raw) < 0.02 else ". ") +
        (f"В топ-10 по важности SHAP на test попали новые признаки: {', '.join(k for k in s4['shap_top'] if k in NEW_FEATURES) or 'нет'}." ), align="j")

    # ---------- 4. модели ----------
    D.h("4. Базовая модель, бустинг и сравнение")
    D.p(f"Квантильный биннинг до 10 корзин с объединением слабых (≥5% наблюдений и ≥5 банкротов в корзине), WoE со сглаживанием; отбор по IV ≥ 0,05 "
        f"и корреляции WoE < 0,7; логистическая регрессия (C = {s2['best_C']} выбрано по TimeSeriesSplit); в карте {s2['n_features']} признаков, "
        f"знаки коэффициентов экономически корректны ({'да' if s2['sign_ok'] else 'нет'}); шкала: 600 баллов при odds 50:1, PDO = 20.",
        bold_prefix="Базовая модель. ", align="j")
    p = s3["final_params"]
    D.p(f"{eng_name}; гиперпараметры подобраны случайным поиском ({s3['n_search']} конфигураций) с оценкой по TimeSeriesSplit внутри train "
        f"(5 расширяющихся окон, обычный KFold не используется), затем ранняя остановка по AUC на validation (итераций: {s3['best_iteration']}). "
        f"Итог: num_leaves = {p['num_leaves']}, learning_rate = {p['learning_rate']}, min_child_samples = {p['min_child_samples']}, "
        f"colsample = {p['colsample_bytree']}, reg_lambda = {p['reg_lambda']}, scale_pos_weight = {p['scale_pos_weight']}; параметры записаны в config.yaml. "
        f"Вероятности откалиброваны isotonic-регрессией на validation (CalibratedClassifierCV, модель заморожена).",
        bold_prefix="Бустинг. ", align="j")
    c = s3["comparison"]
    D.table(["Модель (validation, 15%)", "CV-Gini (TimeSeriesSplit, train)", "Gini", "KS", "PR-AUC", "Brier", "Полнота@10%"],
            [[x["model"], f(x["cv_gini_train"]), f(x["gini"]), f(x["ks"]), f(x["pr_auc"]), f(x["brier"], 4), pc(x["recall@10%"], 0)] for x in c],
            widths=[5.6, 3.0, 1.6, 1.6, 1.8, 1.7, 2.1], size=8)
    d_cv = s3["search_best_cv_gini"] - s2["cv_gini"]
    D.p(f"Сравнение честное: обе модели оценены на одном и том же validation и по одним и тем же окнам TimeSeriesSplit. По CV на train бустинг "
        f"{'лучше' if d_cv > 0 else 'не лучше'} карты на {f(abs(d_cv))} Gini ({f(s3['search_best_cv_gini'])} против {f(s2['cv_gini'])}, "
        f"стд. по окнам ≈ {f(s3['search_best_cv_std'], 2)}). На validation разница {f(s3['valid_raw']['gini'] - s2['valid']['gini'])} "
        f"в пределах шума (95% ДИ Gini бустинга {f(s3['valid_gini_ci'][0])}–{f(s3['valid_gini_ci'][1])}, банкротов {va['bankrupt']}). "
        f"Калибровка не меняет ранжирование, но улучшает Brier ({f(c[1]['brier'], 4)} → {f(c[2]['brier'], 4)}; на validation оценка для isotonic "
        f"оптимистична, т.к. калибратор обучен на этих же данных — честная проверка на test).", align="j")

    # ---------- 5. финансы ----------
    D.h("5. Out-of-time оценка и финансовый эффект")
    D.p(f"Test (последние 15%, {s4['test_n']} компаний, {s4['test_bankrupt']} банкротств) оценён один раз.", bold_prefix="Метрики на test. ")
    D.table(["Модель (test)", "Gini [95% ДИ]", "KS", "PR-AUC", "Полнота@5%", "Полнота@10%", "Полнота@20%"],
            [[nm, f"{f(x['gini'])} [{f(x['gini_ci'][0], 2)}–{f(x['gini_ci'][1], 2)}]", f(x["ks"]), f(x["pr_auc"]),
              pc(x["recall@5%"], 0), pc(x["recall@10%"], 0), pc(x["recall@20%"], 0)]
             for nm, x in (("Скоринговая карта", Bs), (f"Бустинг ({eng_name})", B))],
            widths=[4.2, 3.6, 1.5, 1.7, 2.0, 2.0, 2.0], size=8)
    gd = s4["gini_diff"]
    verdict = ("бустинг значимо лучше карты" if gd["ci"][0] > 0 else ("карта значимо лучше бустинга" if gd["ci"][1] < 0 else
               "различие между моделями статистически незначимо (интервал включает 0)"))
    D.p(f"Разность Gini (бустинг − карта) на test: {f(gd['mean'])}, 95% ДИ [{f(gd['ci'][0])}; {f(gd['ci'][1])}] — {verdict}. "
        f"При проверке {pc(0.10, 0)} компаний с наибольшим скором бустинг ловит {pc(B['recall@10%'], 0)} банкротов (лифт {f(B['lift@10%'], 1)}), "
        f"карта — {pc(Bs['recall@10%'], 0)}. Малое число банкротов в test ({s4['test_bankrupt']}) делает интервалы широкими; различия в "
        f"1–2 банкрота (один банкрот = {f(100 / s4['test_bankrupt'], 1)} п.п. полноты) не являются доказательством превосходства одной из моделей.", align="j")
    cc = fin["calibration_compare"]
    D.p(f"Brier на test: до калибровки {f(cc['raw']['brier'], 4)}, isotonic {f(cc['isotonic']['brier'], 4)}, sigmoid (Platt) {f(cc['sigmoid']['brier'], 4)}; "
        f"средняя PD {pc(cc['isotonic']['mean_pd'], 2)} при фактической доле {pc(cc['actual_rate'], 2)}. Калибратор ступенчатый ({va['bankrupt']} банкротов на "
        f"validation), поэтому PD дискретны, а для PD введена нижняя граница.", bold_prefix="Калибровка. ", align="j")
    m0, lgd = fin["margin"], fin["lgd"]
    D.p(f"Каждая компания — экспозиция EAD = 1. Одобренная «хорошая» компания приносит маржу {pc(m0, 0)} (допущение: в данных нет сумм и ставок), "
        f"одобренная обанкротившаяся приводит к потере LGD × EAD = {pc(lgd, 0)}. Порог PD выбран на validation по максимуму ожидаемой "
        f"прибыли Σ[(1 − PD)·m − PD·LGD] по калиброванным PD (теоретический ориентир m/(m+LGD) = {pc(m0 / (m0 + lgd), 1)}; на практике "
        f"{pc(fin['threshold_pd'], 0)} из-за ступенчатости калибратора); затем на test посчитана реализованная прибыль.", bold_prefix="Деньги. ", align="j")
    rows = []
    for nm, k in (("Одобрить всех", "all_approved"), ("Скоринговая карта + порог", "baseline"), ("Бустинг + порог", "boosting"), ("Идеал (знаем исход)", "oracle")):
        x = fin[k]
        rows.append([nm, pc(x["approve_rate"], 1), x["bad_approved"], f(x["income"], 2), f(x["loss"], 2), f(x["profit"], 2),
                     pc(x["profit_per_exposure"], 2), f(x["profit_per_exposure"] * 1000, 1)])
    D.table(["Стратегия (test)", "Одобрено", "Одобр. банкротов", "Доход", "Потери", "Прибыль", "Прибыль, % EAD", "млн ₸ на 1000 комп. при EAD 1 млн ₸"], rows,
            widths=[3.8, 1.7, 1.9, 1.5, 1.5, 1.6, 2.0, 3.0], size=8)
    gb, gall = fin["boosting"]["profit_per_exposure"], fin["all_approved"]["profit_per_exposure"]
    sens = fin["sensitivity"]
    D.p(f"Бустинг с порогом даёт {pc(gb, 2)} от экспозиции против {pc(gall, 2)} при одобрении всех (+{f((gb - gall) * 100, 2)} п.п.); карта — "
        f"{pc(fin['baseline']['profit_per_exposure'], 2)}; потолок при полном знании исхода — {pc(fin['oracle']['profit_per_exposure'], 2)}. "
        f"К анализу чувствительности: при марже {pc(sens[0]['margin'], 0)}–{pc(sens[-1]['margin'], 0)} выигрыш против «одобрить всех» составляет "
        f"{f(sens[0]['gain_vs_all'] * 100, 2)}–{f(sens[-1]['gain_vs_all'] * 100, 2)} п.п. EAD. Эффект невелик, потому что банкротов мало "
        f"({pc(s4['test_bankrupt'] / s4['test_n'], 1)}), а эффект сильнее при более высоком LGD или меньшей марже.", align="j")
    D.image(FIG / "fig_finance.png", 15.2, "Рис. 2. Реализованная прибыль и кривая полноты на test")

    # ---------- 6. SHAP ----------
    D.h(f"6. Интерпретация: {('SHAP' if s4['shap_method'] == 'shap' else 'вклады признаков')}")
    top = list(s4["shap_top"].items())[:5]
    dirn = s4.get("shap_direction", {})
    def _d(k):
        v = dirn.get(k)
        return "" if v is None else (" — рост значения повышает риск" if v > 0.2 else (" — рост значения снижает риск" if v < -0.2 else " — связь немонотонна"))
    D.p("Главные драйверы решения модели (среднее |вклад| в log-odds на test): " + "; ".join(f"{k} ({f(v, 2)}{_d(k)})" for k, v in top) +
        ". Направление определено по знаку ранговой корреляции значения признака и его вклада.", align="j")
    D.image(FIG / "fig_shap_summary.png", 11.5, "Рис. 3. Summary plot на test (красный — высокое значение признака)")
    D.p("Для трёх наблюдений из test разобраны вклады пяти сильнейших признаков (вклад > 0 повышает PD).", bold_prefix="Разбор наблюдений. ")
    rows = []
    for cs in s4["cases"]:
        drv = "; ".join(f"{t['feature']} = {t['value']:.3g} ({t['shap']:+.2f})" for t in cs["top"][:4])
        rows.append([cs["title"], pc(cs["pd_calibrated"], 1), ("банкрот" if cs["target"] else "не банкрот") + " / " + cs["decision"], drv])
    D.table(["Наблюдение", "PD", "Факт / решение", "Главные признаки: значение (вклад в log-odds)"], rows, widths=[3.8, 1.3, 2.6, 9.3], size=7.5, align_right_from=-1)
    a, b_, c_ = s4["cases"]
    ta, tb = ([t["feature"] for t in x["top"]] for x in (a, b_))
    common = [x for x in ta[:3] if x in tb[:3]]
    up_c = [t for t in c_["top"] if t["shap"] > 0]
    dn_c = [t for t in c_["top"] if t["shap"] < 0]
    fmt = lambda L: ", ".join(f"{t['feature']} ({t['shap']:+.2f})".replace(".", ",") for t in L)
    D.p(f"Наблюдение A (верный отказ): PD {pc(a['pd_calibrated'], 0)}; решающие признаки — {', '.join(ta[:3])}. "
        f"Наблюдение B — ложная тревога (компания не обанкротилась), PD {pc(b_['pd_calibrated'], 0)}; "
        + (f"главные признаки риска те же ({', '.join(common)}): модель видит у двух компаний похожий профиль, и эти признаки не позволили их различить. " if common else
           f"решающие признаки — {', '.join(tb[:3])}. ")
        + f"Наблюдение C — пропущенный банкрот (PD {pc(c_['pd_calibrated'], 1)}): риск повышали {fmt(up_c)}"
        + (f", но частично их гасили признаки, снижающие PD: {fmt(dn_c)}; итоговая PD осталась ниже порога отказа." if dn_c else "; сигнал оказался слабее порога отказа."), align="j")

    # ---------- 7. стабильность ----------
    D.h("7. Стабильность (PSI)")
    psi = s4["psi"]
    D.table(["Показатель", "PSI train→valid", "PSI train→test", "Медиана train", "Медиана test", "Вывод"],
            [[x["item"], f(x["psi_train_valid"], 3), f(x["psi_train_test"], 3), f"{x['train_median']:.3g}".replace(".", ","), f"{x['test_median']:.3g}".replace(".", ","), x["flag"]] for x in psi],
            widths=[5.2, 2.2, 2.2, 2.2, 2.2, 3.2], size=8)
    bad = [x for x in psi if x["psi_train_test"] > 0.25 and not x["item"].startswith("score")]
    score_psi = psi[0]["psi_train_test"]
    txt = (f"PSI скора бустинга train→test равен {f(score_psi)} — {'скор стабилен' if score_psi < 0.1 else ('умеренный сдвиг' if score_psi < 0.25 else 'скор нестабилен')}. ")
    if bad:
        txt += (f"Нестабильны (PSI > 0,25) {len(bad)} из {len(psi) - 2} топ-признаков: " + ", ".join(f"{x['item']} ({f(x['psi_train_test'], 2)})" for x in bad) + ". ")
        near0 = [x for x in bad if abs(x["train_median"]) < 0.01]
        txt += ("Вероятные причины: (а) дрейф во времени — test это более поздние строки, где доля банкротств упала с "
                f"{pc(tr['bankrupt_rate'], 1)} до {pc(te['bankrupt_rate'], 1)}, т.е. изменился состав компаний и условия; "
                + (f"(б) {len(near0)} из них — коэффициенты со значениями у нуля и тяжёлым правым хвостом (медиана train < 0,01): "
                   "небольшой абсолютный сдвиг перекладывает наблюдения между децильными корзинами и раздувает PSI. " if near0 else "")
                + "Рекомендация: мониторить эти признаки при внедрении, переобучать модель по расписанию и сравнивать вклад нестабильных признаков "
                  "в решение; если их вклад велик, заменить их более стабильными аналогами.")
    else:
        txt += "Среди топ-признаков нестабильных (PSI > 0,25) нет."
    D.p(txt, align="j")

    # ---------- 8. ограничения ----------
    D.h("8. Ограничения и выводы")
    D.bullets([
        ("Нет даты. ", "Порядок строк — прокси времени; календарная валидность разбиения подтверждена лишь дрейфом доли банкротств. Результат нельзя считать полноценным out-of-time на календарных данных."),
        ("Мало событий. ", f"В validation {va['bankrupt']}, в test {te['bankrupt']} банкротств: метрики шумные, различия моделей в пределах доверительных интервалов; для надёжности использован TimeSeriesSplit на train и бутстреп-ДИ."),
        ("Допущения денежной оценки. ", f"Маржа {pc(m0, 0)} и EAD = 1 заданы условно (в данных нет сумм); LGD = 45% из задания; важны не абсолютные цифры, а сравнение стратегий."),
        ("Данные. ", "Тайвань, 1999–2009, обезличенные нормированные коэффициенты; перенос на другой рынок (например, эмитентов KASE) потребует переобучения и проверки стабильности."),
        ("Калибровка. ", "Isotonic по " + str(va["bankrupt"]) + " банкротам на validation даёт ступенчатые PD; для продуктивного использования нужна калибровка на большей выборке."),
    ], size=9)
    cmp_txt = ("Бустинг и скоринговая карта сопоставимы по качеству; карта проще и прозрачнее (баллы по каждой корзине), бустинг даёт запас "
               "по гибкости и объяснимость через SHAP" if abs(gd["mean"]) < 0.03 else
               ("Бустинг превосходит карту" if gd["mean"] > 0 else "Карта превосходит бустинг") + " по Gini на test")
    D.p(f"{cmp_txt}. Работа воспроизводима одной командой (python run_all.py), код вынесен в src/, параметры — в config.yaml, "
        f"модель и функция predict(df) — в model.pkl и src/predict.py.", bold_prefix="Итог. ", align="j")

    out = out or (ROOT / "reports" / "report.docx")
    D.d.save(out)
    return str(out)


if __name__ == "__main__":
    print(build_report())
